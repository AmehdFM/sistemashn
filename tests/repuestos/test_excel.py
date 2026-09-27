"""Pruebas de la extensión de columnas de Excel de Repuestos (T2.6b)."""

import pytest
from sqlalchemy import select

from sistemashn.comercial.catalogo.excel import ExcelImportService
from sistemashn.repuestos.excel import RepuestosExcelExtension
from sistemashn.repuestos.partes.models import Part


@pytest.fixture
def excel_service(session_factory, authorizer, clock, catalog_service) -> ExcelImportService:
    return ExcelImportService(
        session_factory,
        authorizer,
        clock=clock,
        catalog_service=catalog_service,
        extensions=[RepuestosExcelExtension()],
    )


def _fila_base(**overrides) -> dict:
    fila = {
        "codigo": "PH-16",
        "codigo_barras": "",
        "nombre": "Filtro de aceite",
        "descripcion": "",
        "categoria": "",
        "unidad": "UND",
        "tasa_isv": 15,
        "precio_venta": 100,
        "stock_minimo": "",
        "activo": "si",
        "numero_parte": "PH-16",
        "fabricante": "Purolator",
        "origen": "original",
    }
    fila.update(overrides)
    return fila


def _escribir_workbook(path, filas) -> None:
    from openpyxl import Workbook

    columnas = list(_fila_base().keys())
    wb = Workbook()
    ws = wb.active
    ws.title = "Productos"
    ws.append(columnas)
    for fila in filas:
        ws.append([fila.get(c, "") for c in columnas])
    wb.save(path)


def test_importar_columnas_de_repuesto_crea_rep_part(
    excel_service, admin_actor, unidad_id, session_factory, tmp_path
) -> None:
    path = tmp_path / "repuestos.xlsx"
    _escribir_workbook(path, [_fila_base()])

    preview = excel_service.preview(admin_actor, path)
    assert preview.errors == []
    resultado = excel_service.commit(admin_actor, preview)
    assert resultado.created == 1

    with session_factory() as session:
        producto = session.execute(select(Part).where(Part.part_number == "PH-16")).scalar_one()
        assert producto.manufacturer == "Purolator"
        assert producto.origin == "original"


def test_origen_vacio_es_generico(excel_service, admin_actor, unidad_id, session_factory, tmp_path):
    path = tmp_path / "repuestos.xlsx"
    _escribir_workbook(path, [_fila_base(origen="")])

    preview = excel_service.preview(admin_actor, path)
    assert preview.errors == []
    excel_service.commit(admin_actor, preview)

    with session_factory() as session:
        parte = session.execute(select(Part).where(Part.part_number == "PH-16")).scalar_one()
        assert parte.origin == "generico"


def test_numero_parte_vacio_no_crea_rep_part(
    excel_service, admin_actor, unidad_id, session_factory, tmp_path
) -> None:
    path = tmp_path / "repuestos.xlsx"
    _escribir_workbook(path, [_fila_base(numero_parte="", codigo="SIN-PARTE")])

    preview = excel_service.preview(admin_actor, path)
    assert preview.errors == []
    excel_service.commit(admin_actor, preview)

    with session_factory() as session:
        partes = session.execute(select(Part)).scalars().all()
        assert partes == []


def test_origen_invalido_es_error_de_fila(excel_service, admin_actor, unidad_id, tmp_path) -> None:
    path = tmp_path / "repuestos.xlsx"
    _escribir_workbook(path, [_fila_base(origen="importado")])

    preview = excel_service.preview(admin_actor, path)
    assert any(e.column.startswith("numero_parte") or e.column == "origen" for e in preview.errors)
    assert preview.rows == []


def test_export_incluye_columnas_de_repuesto(
    excel_service, admin_actor, unidad_id, tmp_path
) -> None:
    path = tmp_path / "repuestos.xlsx"
    _escribir_workbook(path, [_fila_base()])
    preview = excel_service.preview(admin_actor, path)
    excel_service.commit(admin_actor, preview)

    export_path = tmp_path / "export.xlsx"
    excel_service.export_products(admin_actor, export_path)

    from openpyxl import load_workbook

    wb = load_workbook(export_path)
    ws = wb["Productos"]
    encabezados = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert "numero_parte" in encabezados
    assert "fabricante" in encabezados
    assert "origen" in encabezados
    fila = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
    datos = dict(zip(encabezados, fila, strict=True))
    assert datos["numero_parte"] == "PH-16"
    assert datos["fabricante"] == "Purolator"
    assert datos["origen"] == "original"
