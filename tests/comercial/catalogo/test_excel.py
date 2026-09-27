"""Pruebas de importación/exportación de catálogo por Excel (plan T2.5)."""

import time

import pytest
from openpyxl import Workbook, load_workbook
from scripts.generar_catalogo_demo import (
    CATEGORIAS_REQUERIDAS,
    UNIDADES_REQUERIDAS,
    escribir_demo,
    generar_filas,
)

from sistemashn.comercial.catalogo.excel import BASE_COLUMNS, ExcelImportService
from sistemashn.comercial.catalogo.schemas import CategoryInput, UnitInput
from sistemashn.core.errors import PermissionDenied, ValidationError

from ..conftest import make_product_input


@pytest.fixture
def excel_service(session_factory, authorizer, clock, catalog_service) -> ExcelImportService:
    return ExcelImportService(
        session_factory, authorizer, clock=clock, catalog_service=catalog_service
    )


def _crear_categorias_y_unidades(catalog_service, admin_actor) -> None:
    for nombre in CATEGORIAS_REQUERIDAS:
        catalog_service.create_category(admin_actor, CategoryInput(name=nombre))
    for codigo in UNIDADES_REQUERIDAS:
        fraccion = codigo == "LTR"
        catalog_service.create_unit(
            admin_actor, UnitInput(code=codigo, name=codigo.title(), allows_fraction=fraccion)
        )


def _escribir_workbook(path, columnas, filas) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Productos"
    ws.append(columnas)
    for fila in filas:
        ws.append(fila)
    wb.save(path)


def test_plantilla_tiene_encabezados(excel_service, tmp_path) -> None:
    path = tmp_path / "plantilla.xlsx"
    excel_service.write_template(path)

    wb = load_workbook(path)
    assert "Productos" in wb.sheetnames
    assert "Instrucciones" in wb.sheetnames
    encabezados = [c.value for c in next(wb["Productos"].iter_rows(min_row=1, max_row=1))]
    assert encabezados == list(BASE_COLUMNS)


def test_importar_1000_productos_menos_de_10_segundos(
    excel_service, admin_actor, catalog_service, tmp_path
) -> None:
    _crear_categorias_y_unidades(catalog_service, admin_actor)
    path = tmp_path / "demo-1000.xlsx"
    escribir_demo(path, 1000)

    inicio = time.monotonic()
    preview = excel_service.preview(admin_actor, path)
    assert preview.errors == []
    assert len(preview.rows) == 1000
    resultado = excel_service.commit(admin_actor, preview, policy="solo_validas")
    duracion = time.monotonic() - inicio

    assert resultado.created == 1000
    assert resultado.updated == 0
    assert duracion < 10.0


def test_acentos_se_conservan(excel_service, admin_actor, catalog_service, tmp_path) -> None:
    _crear_categorias_y_unidades(catalog_service, admin_actor)
    path = tmp_path / "acentos.xlsx"
    _escribir_workbook(
        path,
        BASE_COLUMNS,
        [("REP-A1", "", "Bandín de distribución", "", "Motor", "UND", 15, 100, 1, "si")],
    )
    preview = excel_service.preview(admin_actor, path)
    assert preview.errors == []
    assert preview.rows[0].data["name"] == "Bandín de distribución"


def test_precio_texto_invalido(excel_service, admin_actor, catalog_service, tmp_path) -> None:
    _crear_categorias_y_unidades(catalog_service, admin_actor)
    path = tmp_path / "precio_malo.xlsx"
    _escribir_workbook(
        path,
        BASE_COLUMNS,
        [("REP-B1", "", "Producto", "", "Motor", "UND", 15, "no-es-numero", 1, "si")],
    )
    preview = excel_service.preview(admin_actor, path)
    assert preview.rows == []
    assert any(e.column == "precio_venta" for e in preview.errors)


def test_columnas_faltantes(excel_service, admin_actor, tmp_path) -> None:
    path = tmp_path / "faltan.xlsx"
    _escribir_workbook(path, ("codigo", "nombre"), [("REP-C1", "Producto")])
    with pytest.raises(ValidationError):
        excel_service.preview(admin_actor, path)


def test_formula_es_rechazada(excel_service, admin_actor, catalog_service, tmp_path) -> None:
    _crear_categorias_y_unidades(catalog_service, admin_actor)
    path = tmp_path / "formula.xlsx"
    _escribir_workbook(
        path,
        BASE_COLUMNS,
        [("REP-D1", "", "=SUMA(A1:A2)", "", "Motor", "UND", 15, 100, 1, "si")],
    )
    preview = excel_service.preview(admin_actor, path)
    assert preview.rows == []
    assert any("fórmula" in e.message for e in preview.errors)


def test_categoria_inexistente(excel_service, admin_actor, catalog_service, tmp_path) -> None:
    for codigo in UNIDADES_REQUERIDAS:
        catalog_service.create_unit(admin_actor, UnitInput(code=codigo, name=codigo))
    path = tmp_path / "cat_mala.xlsx"
    _escribir_workbook(
        path,
        BASE_COLUMNS,
        [("REP-E1", "", "Producto", "", "NoExiste", "UND", 15, 100, 1, "si")],
    )
    preview = excel_service.preview(admin_actor, path)
    assert preview.rows == []
    assert any(e.column == "categoria" for e in preview.errors)


def test_codigo_duplicado_en_archivo(excel_service, admin_actor, catalog_service, tmp_path) -> None:
    _crear_categorias_y_unidades(catalog_service, admin_actor)
    path = tmp_path / "dup.xlsx"
    _escribir_workbook(
        path,
        BASE_COLUMNS,
        [
            ("REP-F1", "", "Producto uno", "", "Motor", "UND", 15, 100, 1, "si"),
            ("REP-F1", "", "Producto uno repetido", "", "Motor", "UND", 15, 100, 1, "si"),
        ],
    )
    preview = excel_service.preview(admin_actor, path)
    assert len(preview.rows) == 1
    assert any(e.row_number == 3 and e.column == "codigo" for e in preview.errors)


def test_politica_todo_o_nada_no_escribe_nada_con_un_error(
    excel_service, admin_actor, catalog_service, session_factory, tmp_path
) -> None:
    from sistemashn.comercial.catalogo.models import Product

    _crear_categorias_y_unidades(catalog_service, admin_actor)
    path = tmp_path / "mixto.xlsx"
    _escribir_workbook(
        path,
        BASE_COLUMNS,
        [
            ("REP-G1", "", "Producto válido", "", "Motor", "UND", 15, 100, 1, "si"),
            ("REP-G2", "", "Producto inválido", "", "NoExiste", "UND", 15, 100, 1, "si"),
        ],
    )
    preview = excel_service.preview(admin_actor, path)
    assert len(preview.rows) == 1
    assert len(preview.errors) == 1

    with pytest.raises(ValidationError):
        excel_service.commit(admin_actor, preview, policy="todo_o_nada")

    with session_factory() as session:
        assert session.query(Product).count() == 0


def test_politica_solo_validas_escribe_las_validas(
    excel_service, admin_actor, catalog_service, tmp_path
) -> None:
    _crear_categorias_y_unidades(catalog_service, admin_actor)
    path = tmp_path / "mixto2.xlsx"
    _escribir_workbook(
        path,
        BASE_COLUMNS,
        [
            ("REP-H1", "", "Producto válido", "", "Motor", "UND", 15, 100, 1, "si"),
            ("REP-H2", "", "Producto inválido", "", "NoExiste", "UND", 15, 100, 1, "si"),
        ],
    )
    preview = excel_service.preview(admin_actor, path)
    resultado = excel_service.commit(admin_actor, preview, policy="solo_validas")
    assert resultado.created == 1
    assert resultado.skipped == 1


def test_reimportacion_es_idempotente(
    excel_service, admin_actor, catalog_service, tmp_path
) -> None:
    _crear_categorias_y_unidades(catalog_service, admin_actor)
    path = tmp_path / "reimport.xlsx"
    escribir_demo(path, 50)

    preview1 = excel_service.preview(admin_actor, path)
    resultado1 = excel_service.commit(admin_actor, preview1, policy="solo_validas")
    assert resultado1.created == 50
    assert resultado1.updated == 0

    preview2 = excel_service.preview(admin_actor, path)
    resultado2 = excel_service.commit(admin_actor, preview2, policy="solo_validas")
    assert resultado2.created == 0
    assert resultado2.updated == 50

    with excel_service.factory() as session:
        from sistemashn.comercial.catalogo.models import Product

        assert session.query(Product).count() == 50


def test_export_import_no_cambia_productos(
    excel_service, admin_actor, catalog_service, unidad_id, categoria_id, tmp_path
) -> None:
    catalog_service.create_product(admin_actor, make_product_input(unidad_id, categoria_id))
    catalog_service.create_product(
        admin_actor,
        make_product_input(unidad_id, categoria_id, code="ABC-999", name="Otro producto"),
    )

    path = tmp_path / "export.xlsx"
    excel_service.export_products(admin_actor, path)

    preview = excel_service.preview(admin_actor, path)
    assert preview.errors == []
    resultado = excel_service.commit(admin_actor, preview, policy="solo_validas")
    assert resultado.created == 0
    assert resultado.updated == 2


def test_archivo_corrupto(excel_service, admin_actor, tmp_path) -> None:
    path = tmp_path / "corrupto.xlsx"
    path.write_bytes(b"no soy un archivo de excel valido, solo bytes al azar 1234567890")
    with pytest.raises(ValidationError):
        excel_service.preview(admin_actor, path)


def test_vendedor_no_puede_importar(excel_service, vendedor_actor, tmp_path) -> None:
    path = tmp_path / "vendedor.xlsx"
    _escribir_workbook(path, BASE_COLUMNS, [])
    with pytest.raises(PermissionDenied):
        excel_service.preview(vendedor_actor, path)


def test_generar_filas_produce_datos_variados() -> None:
    filas = generar_filas(1000)
    assert len(filas) == 1000
    codigos = {f["codigo"] for f in filas}
    assert len(codigos) == 1000
    assert any("í" in f["nombre"] or "ñ" in f["nombre"] or "ó" in f["nombre"] for f in filas)
