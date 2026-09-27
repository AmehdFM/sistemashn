"""Importación y exportación de catálogo por Excel (plan T2.5)."""

from collections.abc import Callable
from dataclasses import dataclass
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Literal, Protocol
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError as PydanticValidationError
from sqlalchemy import and_, select
from sqlalchemy.orm import Session, sessionmaker

from sistemashn.comercial.catalogo.models import Category, Product, Unit
from sistemashn.comercial.catalogo.schemas import ProductInput
from sistemashn.comercial.catalogo.search import normalize_search
from sistemashn.comercial.catalogo.service import CatalogService
from sistemashn.comercial.inventario.models import Stock
from sistemashn.core.audit.service import audit
from sistemashn.core.authorization.actor import Actor
from sistemashn.core.authorization.service import Authorizer
from sistemashn.core.db.uow import run_in_transaction
from sistemashn.core.errors import ValidationError

BASE_COLUMNS: tuple[str, ...] = (
    "codigo",
    "codigo_barras",
    "nombre",
    "descripcion",
    "categoria",
    "unidad",
    "tasa_isv",
    "precio_venta",
    "stock_minimo",
    "activo",
)

_SHEET_PRODUCTOS = "Productos"
_TASAS_VALIDAS = (Decimal("0"), Decimal("15"), Decimal("18"))
_VALORES_SI = {"si", "sí", "yes", "true", "1"}
_VALORES_NO = {"no", "false", "0"}


def _utcnow() -> datetime:
    return datetime.now(UTC)


def _to_decimal(value: object) -> Decimal:
    """Convierte una celda numérica a `Decimal` sin pasar por `float`."""
    if isinstance(value, Decimal):
        return value
    if isinstance(value, bool):
        raise ValueError("se esperaba un número")
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        return Decimal(str(value))
    if value is None:
        raise ValueError("valor numérico requerido")
    texto = str(value).strip()
    if not texto:
        raise ValueError("valor numérico requerido")
    return Decimal(texto)


class ExcelColumnExtension(Protocol):
    """Hook para que otros módulos (p. ej. Repuestos) agreguen columnas al Excel."""

    columns: tuple[str, ...]

    def validate(self, row: dict[str, Any]) -> list[str]:
        """Devuelve mensajes de error para las columnas de la extensión en `row`."""
        ...

    def apply(self, session: Session, actor: Actor, product_id: int, row: dict[str, Any]) -> None:
        """Aplica los valores de la extensión al producto, dentro de la transacción."""
        ...

    def export(self, session: Session, product_id: int) -> dict[str, Any]:
        """Devuelve los valores de la extensión para exportar el producto."""
        ...


@dataclass(frozen=True)
class RowError:
    row_number: int
    column: str
    message: str


@dataclass(frozen=True)
class PreviewRow:
    row_number: int
    codigo: str
    action: Literal["crear", "actualizar"]
    data: dict[str, Any]


@dataclass(frozen=True)
class ImportPreview:
    rows: list[PreviewRow]
    errors: list[RowError]
    total_rows: int


@dataclass(frozen=True)
class ImportResult:
    created: int
    updated: int
    skipped: int


def _fila_vacia(valores: dict[str, Any]) -> bool:
    for valor in valores.values():
        if valor is None:
            continue
        if isinstance(valor, str) and not valor.strip():
            continue
        return False
    return True


class ExcelImportService:
    """Importa/exporta el catálogo desde/hacia archivos Excel (`.xlsx`)."""

    def __init__(
        self,
        factory: sessionmaker[Session],
        authorizer: Authorizer,
        clock: Callable[[], datetime] = _utcnow,
        catalog_service: CatalogService | None = None,
        extensions: list[ExcelColumnExtension] | None = None,
    ) -> None:
        self.factory = factory
        self.authorizer = authorizer
        self.clock = clock
        self.catalog_service = catalog_service
        self.extensions = extensions or []

    # -- Plantilla ----------------------------------------------------------

    def write_template(self, path: str | Path) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = _SHEET_PRODUCTOS
        columnas = list(BASE_COLUMNS) + [c for ext in self.extensions for c in ext.columns]
        ws.append(columnas)
        for cell in ws[1]:
            cell.font = Font(bold=True)

        instrucciones = wb.create_sheet("Instrucciones")
        instrucciones.append(("Columna", "Descripción"))
        for cell in instrucciones[1]:
            cell.font = Font(bold=True)
        explicaciones = [
            ("codigo", "Código único del producto. Obligatorio, hasta 40 caracteres."),
            ("codigo_barras", "Código de barras. Opcional, debe ser único si se indica."),
            ("nombre", "Nombre del producto. Obligatorio, hasta 200 caracteres."),
            ("descripcion", "Descripción libre. Opcional."),
            (
                "categoria",
                "Nombre de una categoría ya existente. Opcional; si se indica, debe existir.",
            ),
            ("unidad", "Código de una unidad ya existente. Obligatorio (por ejemplo UND o LTR)."),
            ("tasa_isv", "Tasa de ISV: 0, 15 o 18."),
            ("precio_venta", "Precio de venta, número con hasta 2 decimales."),
            (
                "stock_minimo",
                "Cantidad mínima antes de alertar stock bajo. Opcional, 0 por defecto.",
            ),
            ("activo", "si o no. Vacío equivale a si."),
        ]
        for ext in self.extensions:
            for columna in ext.columns:
                explicaciones.append((columna, "Columna adicional registrada por otro módulo."))
        for fila in explicaciones:
            instrucciones.append(fila)

        wb.save(path)

    # -- Vista previa ---------------------------------------------------------

    def preview(self, actor: Actor, path: str | Path) -> ImportPreview:
        def _op(session: Session) -> ImportPreview:
            self.authorizer.require(session, actor, "com.catalogo.importar")
            return self._build_preview(session, path)

        return run_in_transaction(self.factory, _op, readonly=True)

    def _load_workbook(self, path: str | Path):  # type: ignore[no-untyped-def]
        try:
            return load_workbook(path, data_only=False, read_only=True)
        except (BadZipFile, InvalidFileException, KeyError, OSError) as exc:
            raise ValidationError("archivo de Excel inválido o dañado") from exc

    def _productos_sheet(self, wb):  # type: ignore[no-untyped-def]
        if _SHEET_PRODUCTOS not in wb.sheetnames:
            raise ValidationError(f"falta la hoja '{_SHEET_PRODUCTOS}'")
        return wb[_SHEET_PRODUCTOS]

    def _header_index(self, ws) -> dict[str, int]:  # type: ignore[no-untyped-def]
        primera = next(iter(ws.iter_rows(min_row=1, max_row=1)), ())
        header: dict[str, int] = {}
        for idx, cell in enumerate(primera):
            valor = cell.value
            if valor is None:
                continue
            header[str(valor).strip().lower()] = idx
        faltantes = [c for c in BASE_COLUMNS if c not in header]
        if faltantes:
            raise ValidationError("faltan columnas obligatorias: " + ", ".join(faltantes))
        return header

    def _build_preview(self, session: Session, path: str | Path) -> ImportPreview:
        wb = self._load_workbook(path)
        ws = self._productos_sheet(wb)
        header_index = self._header_index(ws)
        columnas_extension = [c for ext in self.extensions for c in ext.columns]
        columnas_conocidas = list(BASE_COLUMNS) + columnas_extension

        categorias = {c.name: c.id for c in session.scalars(select(Category)).all()}
        unidades = {u.code: u.id for u in session.scalars(select(Unit)).all()}
        productos_existentes = {p.code for p in session.scalars(select(Product)).all()}

        rows: list[PreviewRow] = []
        errors: list[RowError] = []
        seen_codes: set[str] = set()
        total_rows = 0

        for row_number, cells in enumerate(ws.iter_rows(min_row=2), start=2):
            valores = {
                col: cells[idx].value if idx < len(cells) else None
                for col, idx in header_index.items()
                if col in columnas_conocidas
            }
            if _fila_vacia(valores):
                continue
            total_rows += 1

            formula_cols = []
            for col in columnas_conocidas:
                idx = header_index.get(col)
                if idx is None or idx >= len(cells):
                    continue
                cell = cells[idx]
                valor = cell.value
                if getattr(cell, "data_type", None) == "f" or (
                    isinstance(valor, str) and valor.startswith("=")
                ):
                    formula_cols.append(col)
            if formula_cols:
                for col in formula_cols:
                    errors.append(
                        RowError(row_number, col, "la celda contiene una fórmula, no un valor")
                    )
                continue

            raw = {col: valores.get(col) for col in columnas_conocidas}
            row_errors: list[RowError] = []

            codigo = str(raw.get("codigo") or "").strip().upper()
            if codigo:
                if codigo in seen_codes:
                    row_errors.append(
                        RowError(row_number, "codigo", f"código duplicado en el archivo: {codigo}")
                    )
                seen_codes.add(codigo)

            category_id = None
            categoria_nombre = raw.get("categoria")
            if categoria_nombre not in (None, ""):
                categoria_nombre = str(categoria_nombre).strip()
                category_id = categorias.get(categoria_nombre)
                if category_id is None:
                    row_errors.append(
                        RowError(
                            row_number, "categoria", f"categoría inexistente: '{categoria_nombre}'"
                        )
                    )

            unit_id = None
            unidad_valor = raw.get("unidad")
            unidad_codigo = (
                str(unidad_valor).strip().upper() if unidad_valor not in (None, "") else ""
            )
            if unidad_codigo:
                unit_id = unidades.get(unidad_codigo)
            if unit_id is None:
                row_errors.append(
                    RowError(row_number, "unidad", f"unidad inexistente: '{unidad_valor}'")
                )

            tax_rate: Decimal | None = None
            try:
                tasa_raw = _to_decimal(raw.get("tasa_isv"))
                if tasa_raw not in _TASAS_VALIDAS:
                    raise ValueError
                tax_rate = (tasa_raw / Decimal("100")).quantize(Decimal("0.01"))
            except Exception:
                row_errors.append(
                    RowError(row_number, "tasa_isv", "tasa de ISV inválida: debe ser 0, 15 o 18")
                )

            precio: Decimal | None = None
            try:
                precio = _to_decimal(raw.get("precio_venta")).quantize(Decimal("0.01"))
            except Exception:
                row_errors.append(RowError(row_number, "precio_venta", "precio de venta inválido"))

            stock_min = Decimal("0")
            if raw.get("stock_minimo") not in (None, ""):
                try:
                    stock_min = _to_decimal(raw.get("stock_minimo"))
                except Exception:
                    row_errors.append(RowError(row_number, "stock_minimo", "stock mínimo inválido"))

            activo = True
            activo_raw = raw.get("activo")
            if activo_raw not in (None, ""):
                texto = str(activo_raw).strip().lower()
                if texto in _VALORES_SI:
                    activo = True
                elif texto in _VALORES_NO:
                    activo = False
                else:
                    row_errors.append(
                        RowError(row_number, "activo", "valor de 'activo' inválido: use si/no")
                    )

            barcode = raw.get("codigo_barras")
            barcode = str(barcode).strip() or None if barcode not in (None, "") else None
            nombre = raw.get("nombre")
            nombre = str(nombre).strip() if nombre is not None else ""
            descripcion = raw.get("descripcion")
            descripcion = str(descripcion).strip() if descripcion not in (None, "") else None

            producto_input: ProductInput | None = None
            try:
                producto_input = ProductInput(
                    code=codigo,
                    barcode=barcode,
                    name=nombre,
                    description=descripcion,
                    category_id=category_id,
                    unit_id=unit_id if unit_id is not None else 0,
                    tax_rate=tax_rate if tax_rate is not None else Decimal("0"),
                    sale_price=precio if precio is not None else Decimal("0"),
                    min_stock=stock_min,
                    active=activo,
                )
            except PydanticValidationError as exc:
                for err in exc.errors():
                    columna = str(err["loc"][0]) if err["loc"] else "general"
                    row_errors.append(RowError(row_number, columna, str(err["msg"])))

            if row_errors:
                errors.extend(row_errors)
                continue

            for ext in self.extensions:
                mensajes = ext.validate(raw)
                columna_ext = ext.columns[0] if ext.columns else "extension"
                for mensaje in mensajes:
                    row_errors.append(RowError(row_number, columna_ext, mensaje))

            if row_errors:
                errors.extend(row_errors)
                continue

            assert producto_input is not None
            action: Literal["crear", "actualizar"] = (
                "actualizar" if codigo in productos_existentes else "crear"
            )
            data = {
                "code": producto_input.code,
                "barcode": producto_input.barcode,
                "name": producto_input.name,
                "description": producto_input.description,
                "category_id": producto_input.category_id,
                "unit_id": producto_input.unit_id,
                "tax_rate": producto_input.tax_rate,
                "sale_price": producto_input.sale_price,
                "min_stock": producto_input.min_stock,
                "active": producto_input.active,
                "raw": raw,
            }
            rows.append(PreviewRow(row_number=row_number, codigo=codigo, action=action, data=data))

        return ImportPreview(rows=rows, errors=errors, total_rows=total_rows)

    # -- Confirmación ---------------------------------------------------------

    def commit(
        self,
        actor: Actor,
        preview: ImportPreview,
        policy: Literal["solo_validas", "todo_o_nada"] = "solo_validas",
    ) -> ImportResult:
        def _op(session: Session) -> ImportResult:
            self.authorizer.require(session, actor, "com.catalogo.importar")
            if policy == "todo_o_nada" and preview.errors:
                raise ValidationError(
                    "el archivo tiene errores; no se importó nada (política todo_o_nada): "
                    f"{len(preview.errors)} error(es)"
                )

            ahora = self.clock()
            creados = 0
            actualizados = 0

            for fila in preview.rows:
                datos = fila.data
                producto = session.scalar(select(Product).where(Product.code == fila.codigo))
                self._validar_barcode_unico(
                    session, datos["barcode"], product_id=producto.id if producto else None
                )

                if producto is None:
                    producto = Product(
                        code=datos["code"],
                        barcode=datos["barcode"],
                        name=datos["name"],
                        name_search=normalize_search(datos["name"]),
                        description=datos["description"],
                        category_id=datos["category_id"],
                        unit_id=datos["unit_id"],
                        tax_rate=datos["tax_rate"],
                        sale_price=datos["sale_price"],
                        min_stock=datos["min_stock"],
                        is_kit=False,
                        active=datos["active"],
                        created_at=ahora,
                        updated_at=ahora,
                    )
                    session.add(producto)
                    session.flush()
                    # Igual que `CatalogService.create_product`: todo producto no-kit
                    # nace con su fila de existencias en cero (importar no ajusta stock).
                    session.add(
                        Stock(
                            product_id=producto.id,
                            on_hand=Decimal("0"),
                            reserved=Decimal("0"),
                            unsellable=Decimal("0"),
                            avg_cost=Decimal("0"),
                            updated_at=ahora,
                        )
                    )
                    session.flush()
                    creados += 1
                else:
                    producto.barcode = datos["barcode"]
                    producto.name = datos["name"]
                    producto.name_search = normalize_search(datos["name"])
                    producto.description = datos["description"]
                    producto.category_id = datos["category_id"]
                    producto.unit_id = datos["unit_id"]
                    producto.tax_rate = datos["tax_rate"]
                    producto.sale_price = datos["sale_price"]
                    producto.min_stock = datos["min_stock"]
                    producto.active = datos["active"]
                    producto.updated_at = ahora
                    session.flush()
                    actualizados += 1

                for ext in self.extensions:
                    ext.apply(session, actor, producto.id, datos["raw"])

            omitidos = len({error.row_number for error in preview.errors})

            audit(
                session,
                actor,
                "com.catalogo.importado",
                entity_type="com_product",
                summary=f"Importación de catálogo: {creados} creados, {actualizados} actualizados",
                detail={"creados": creados, "actualizados": actualizados, "omitidos": omitidos},
                clock=self.clock,
            )
            return ImportResult(created=creados, updated=actualizados, skipped=omitidos)

        return run_in_transaction(self.factory, _op)

    def _validar_barcode_unico(
        self, session: Session, barcode: str | None, *, product_id: int | None
    ) -> None:
        if barcode is None:
            return
        condiciones = [Product.barcode == barcode]
        if product_id is not None:
            condiciones.append(Product.id != product_id)
        if session.scalar(select(Product).where(and_(*condiciones))) is not None:
            raise ValidationError(f"el código de barras '{barcode}' ya existe")

    # -- Exportación ---------------------------------------------------------

    def export_products(self, actor: Actor, path: str | Path) -> None:
        def _op(session: Session) -> None:
            self.authorizer.require(session, actor, "com.catalogo.ver")
            categorias = {c.id: c.name for c in session.scalars(select(Category)).all()}
            unidades = {u.id: u.code for u in session.scalars(select(Unit)).all()}
            productos = session.scalars(select(Product).order_by(Product.code)).all()

            wb = Workbook()
            ws = wb.active
            ws.title = _SHEET_PRODUCTOS
            columnas = list(BASE_COLUMNS) + [c for ext in self.extensions for c in ext.columns]
            ws.append(columnas)
            for cell in ws[1]:
                cell.font = Font(bold=True)

            for producto in productos:
                fila: list[Any] = [
                    producto.code,
                    producto.barcode or "",
                    producto.name,
                    producto.description or "",
                    categorias.get(producto.category_id, "") if producto.category_id else "",
                    unidades.get(producto.unit_id, ""),
                    float((producto.tax_rate * Decimal("100")).quantize(Decimal("1"))),
                    float(producto.sale_price.quantize(Decimal("0.01"))),
                    float(producto.min_stock),
                    "si" if producto.active else "no",
                ]
                for ext in self.extensions:
                    datos_extra = ext.export(session, producto.id)
                    for columna in ext.columns:
                        fila.append(datos_extra.get(columna, ""))
                ws.append(fila)

            wb.save(path)

        run_in_transaction(self.factory, _op, readonly=True)
